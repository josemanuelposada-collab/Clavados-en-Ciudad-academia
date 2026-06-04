from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "archivos" / "sustentacion"
OUT_DIR.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT_DIR / "Explicacion_Nivel_1_Fisicas.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    return p


def add_note(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EAF3F8")
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 77, 120)
    p.add_run(body)
    doc.add_paragraph()


def add_range_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_widths(table, [0.9, 1.55, 4.05])
    headers = ["Lineas", "Bloque", "Explicacion para sustentar"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(31, 77, 120)
    for lineas, bloque, explicacion in rows:
        cells = table.add_row().cells
        cells[0].text = lineas
        cells[1].text = bloque
        cells[2].text = explicacion
        for cell in cells:
            set_cell_margins(cell)
    doc.add_paragraph()


def add_formula_table(doc):
    rows = [
        (
            "Gravedad / parabola",
            "ay = gravedad * factorGravedad; vy += ay * dt; y += vy * dt",
            "La velocidad vertical cambia con el tiempo. Por eso no es movimiento rectilineo simple. El salto inicia con vy negativa y luego la gravedad la vuelve positiva.",
        ),
        (
            "Viento turbulento",
            "viento = vientoBase + sin(tiempo * 4 + y * 0.01) * 26",
            "El viento tiene una parte constante por dificultad y una variacion sinusoidal que simula rafagas.",
        ),
        (
            "Piscina con aceleracion",
            "a = sin(tiempo * 0.82) * 34 * factor; v += a * dt; x += v * dt",
            "La piscina no se mueve linealmente: acelera y frena, rebota en limites y aumenta dificultad.",
        ),
        (
            "Campo radial Mikoto",
            "distancia = sqrt(dx^2 + dy^2); fuerza depende de radio e influencia",
            "Las monedas aceleran hacia el jugador segun distancia y radio del poder.",
        ),
        (
            "Control vectorial Accelerator",
            "vx/vy se mezclan con una velocidad objetivo hacia el personaje",
            "No es solo atraer: ajusta directamente el vector velocidad de la moneda.",
        ),
        (
            "Meltdowner Mugino",
            "afecta monedas dentro de un corredor vertical y las canaliza lateralmente",
            "Representa un haz fisico direccional, diferente al campo radial.",
        ),
        (
            "Dark Matter",
            "fuerza radial + fuerza tangencial perpendicular",
            "Produce orbita: la moneda no va recta al jugador, curva alrededor del campo.",
        ),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_widths(table, [1.55, 2.45, 2.5])
    for i, h in enumerate(["Fisica", "Ecuacion / regla", "Como defenderla"]):
        cell = table.cell(0, i)
        set_cell_shading(cell, "E8EEF5")
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(31, 77, 120)
    for fisica, ecuacion, defensa in rows:
        cells = table.add_row().cells
        cells[0].text = fisica
        cells[1].text = ecuacion
        cells[2].text = defensa
        for cell in cells:
            set_cell_margins(cell)
    doc.add_paragraph()


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18
    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Clavados en Ciudad Academia")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(31, 77, 120)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = subtitle.add_run("Explicacion linea a linea del Nivel 1 y sus fisicas")
    rr.font.size = Pt(12)
    rr.italic = True

    add_note(
        doc,
        "Objetivo de sustentacion",
        "Explicar que el Nivel 1 combina POO, memoria dinamica, fisicas por dt, dificultad parametrizada, poderes fisicos y agente inteligente.",
    )

    add_heading(doc, "Resumen rapido para decir al profesor", 1)
    bullets = [
        "El Nivel 1 es una caida desde plataforma hacia una piscina movil.",
        "La posicion del jugador no se cambia a mano: se actualiza con velocidad y aceleracion.",
        "Uso dt para que la simulacion dependa del tiempo entre frames.",
        "La gravedad, el viento, la piscina acelerada y los campos de monedas son fisicas defendibles.",
        "La memoria dinamica aparece con make_unique y contenedores de unique_ptr.",
        "El dron percibe, razona, actua y aprende del error de entrada.",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(b)

    add_heading(doc, "Fisicas que debes sustentar", 1)
    add_formula_table(doc)

    add_heading(doc, "Explicacion linea a linea por bloques", 1)
    add_note(
        doc,
        "Como usar esta tabla",
        "Cuando te pidan una linea exacta, ubica el rango. Si preguntan por fisica, ve directo a los rangos 135-174, 195-254 y 650-737.",
    )

    rows = [
        ("1-5", "Includes", "Incluyen la cabecera del nivel, cache de sprites, algoritmos como min/max/clamp, funciones matematicas como sin/sqrt y gradientes de Qt."),
        ("7-18", "Panel HUD", "Funcion local para dibujar paneles oscuros del HUD. No es fisica, es capa grafica."),
        ("20-28", "Barras HUD", "Dibuja barras de progreso. La linea 22 limita el porcentaje entre 0 y 1 para evitar barras desbordadas."),
        ("30-33", "Hitbox justa", "Contrae un rectangulo de colision. Sirve para que el sprite no choque por sus esquinas transparentes."),
        ("36-40", "Constructor", "Crea jugador, plataforma y dron con make_unique. Esto sustenta memoria dinamica sin fugas."),
        ("41-46", "Mundo fisico", "Define altura del mundo, camara, piscina, zona meta, zona de viento y suelo como rectangulos del espacio del juego."),
        ("48-71", "Recursos", "Carga imagenes del nivel. Si un recurso no carga, el juego todavia puede seguir con alternativas."),
        ("73-92", "Estado inicial", "Inicializa puntaje, intentos, banderas de victoria/derrota, viento, gravedad, energia del poder y variables de piscina."),
        ("94-99", "Dificultad y monedas", "Configura dificultad normal, aplica sus parametros, ubica plataforma, toma intentos maximos y crea monedas."),
        ("101-104", "Posicion inicial", "Coloca al jugador sobre la plataforma calculando el centro horizontal y la altura correcta."),
        ("107", "Destructor", "Destructor por defecto. Como se usan unique_ptr, no se necesita liberar manualmente."),
        ("109-126", "Parametros de dificultad", "Toma viento lateral, energia y gravedad segun dificultad. Facil 30, normal 38, dificil 48."),
        ("128-133", "Cambio de dificultad", "Reconfigura dificultad, aplica parametros y reinicia el nivel."),
        ("135-146", "Inicio del update", "Actualiza plataforma, dron, tiempo, cooldown del poder, piscina movil y monedas cada frame."),
        ("148-153", "Jugador en plataforma", "Si el jugador no esta en el aire, se mantiene pegado a la plataforma."),
        ("155-159", "Gravedad", "Si el intento sigue vivo, verifica viento y aplica gravedad al jugador. Esta es fisica clave."),
        ("160-163", "Turbulencia", "Si esta en zona de viento, calcula una rafaga con sin(tiempo + posicion) y la aplica como aceleracion horizontal."),
        ("165-167", "Interaccion dron-jugador", "Si el dron toca al jugador, aplica viento lateral de empuje. No cambia directamente la posicion."),
        ("169-173", "Integracion y camara", "Actualiza al jugador con velocidades/aceleraciones, revisa colisiones y mueve la camara."),
        ("176-193", "Crear monedas", "Limpia el vector, reserva memoria y crea monedas con make_unique en posiciones del mundo."),
        ("195-202", "Distancia a monedas", "Para cada moneda calcula dx, dy y distancia euclidiana con sqrt(dx^2 + dy^2)."),
        ("204-217", "Poderes por personaje", "Segun personaje aplica modelo fisico distinto: vectorial, haz, orbita o atraccion radial."),
        ("220-225", "Recolectar moneda", "Usa hitbox contraida. Si intersecta, recolecta, suma puntaje y dispara sonido."),
        ("229-233", "Piscina bloqueada", "Si el intento termino o el nivel acabo, la piscina deja de actualizarse."),
        ("235-240", "Piscina acelerada", "Calcula aceleracion sinusoidal, integra velocidad y traslada la piscina. Es fisica defendible."),
        ("242-253", "Rebote y zona meta", "Limita la piscina entre dos bordes y cambia direccion como rebote amortiguado."),
        ("256-259", "Camara", "Usa clamp para seguir al jugador sin salirse del mundo."),
        ("261-264", "Colisiones", "Agrupa verificaciones de piscina y suelo."),
        ("267-273", "Zona de viento", "Detecta si el jugador esta dentro de la zona fisica de viento."),
        ("276-308", "Entrada al agua", "Si el jugador toca piscina mientras cae, calcula puntaje, determina victoria, detiene movimiento y el dron aprende."),
        ("311-329", "Fallo por suelo", "Si toca suelo fuera de la piscina, puntaje cero, pierde intento y se detiene."),
        ("332-379", "Puntaje fisico", "Evalua error respecto al centro de la piscina, velocidad vertical de entrada y monedas recolectadas."),
        ("381-398", "Fondo", "Dibuja fondo vertical del mundo. Si no carga, usa gradiente alternativo."),
        ("400-418", "Viento y piso", "Dibuja rafagas y piso final, haciendo visible la zona fisica."),
        ("420-437", "Piscina visual", "Dibuja base, decoracion y sprite de piscina en la posicion fisica actual."),
        ("439-448", "Entidades visibles", "Dibuja monedas, plataforma, dron y efecto de poder si esta activo."),
        ("450-482", "Entrada visual al agua", "Si termino dentro de piscina, dibuja burbujas, splash segun puntaje y personaje semisumergido."),
        ("485-508", "Rafagas visibles", "Anima viento con seno; conecta visualmente con la turbulencia aplicada al jugador."),
        ("510-571", "Efectos de poder", "Dibuja efecto distinto por personaje: vectorial, Meltdowner, materia oscura o electromagnetismo."),
        ("573-638", "HUD", "Muestra energia, puntaje, monedas, intentos, estado del viento y resultado del intento."),
        ("640-648", "Render general", "Mueve el sistema de coordenadas por camaraY, dibuja mundo y luego HUD fijo."),
        ("650-658", "Salto inicial", "Espacio activa caida: setVX(130) y setVY(-86). Luego gravedad forma trayectoria parabolica."),
        ("661-667", "Control lateral", "A y D cambian estado de movimiento lateral; el Personaje convierte eso en aceleracion horizontal."),
        ("669-692", "Acciones extra", "E activa poder, R reinicia, 1/2/3 cambian dificultad."),
        ("695-704", "Liberar teclas", "Al soltar A o D se desactiva el control lateral."),
        ("706-710", "Mouse", "Click tambien activa el poder."),
        ("712-737", "Activar poder", "Controla duracion, cooldown y radio segun personaje. No es poder infinito."),
        ("739-769", "Reiniciar intento", "Restaura puntaje parcial, monedas, piscina, dron y estado fisico del jugador."),
        ("771-807", "Reiniciar nivel", "Restaura todo el nivel: camara, tiempo, intentos, piscina, energia, plataforma y monedas."),
        ("809-814", "Elegir personaje", "Configura el tipo de personaje, cambia radio del poder y reinicia."),
        ("816-840", "Consultas finales", "Metodos simples para puntaje, estado del nivel, nombre y consumo de sonidos."),
    ]
    add_range_table(doc, rows)

    add_heading(doc, "Guion oral de 60 segundos", 1)
    p = doc.add_paragraph()
    p.add_run(
        "En el Nivel 1 el jugador salta desde una plataforma hacia una piscina movil. "
        "La fisica principal es una integracion por dt: primero se acumulan aceleraciones, luego se actualizan velocidades y por ultimo posiciones. "
        "La gravedad produce una trayectoria parabolica porque el jugador inicia con una velocidad vertical negativa y luego la aceleracion de gravedad cambia esa velocidad. "
        "Tambien hay una zona de viento con turbulencia sinusoidal, una piscina con aceleracion variable y poderes que modifican la velocidad de las monedas. "
        "La dificultad cambia parametros fisicos como gravedad, viento, energia e intentos, y el dron funciona como agente inteligente porque percibe, razona, actua y aprende del error de entrada."
    )

    add_heading(doc, "Preguntas duras y respuestas", 1)
    qa = [
        ("Por que esto no es movimiento rectilineo simple?", "Porque hay aceleracion. La velocidad vertical y horizontal cambian con gravedad, viento y campos."),
        ("Donde esta dt?", "En actualizar(dt). dt escala velocidad, aceleracion, cooldowns y movimiento para no depender directamente de FPS."),
        ("Cual es la fisica mas clara?", "La gravedad: ay = gravedad * factor, vy += ay*dt, y += vy*dt."),
        ("Que hace la dificultad?", "Modifica viento, energia, intentos, puntaje minimo, gravedad y factores de movimiento."),
        ("Como justificas memoria dinamica?", "Con unique_ptr y make_unique para jugador, plataforma, dron y monedas, evitando delete manual."),
        ("Como justificas POO?", "Entidad define actualizar/dibujar; Personaje, Moneda, Plataforma y Dron especializan comportamiento."),
    ]
    add_range_table(doc, [(q, "Respuesta", a) for q, a in qa])

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_doc())
